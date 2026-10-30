import io
from datetime import date
from decimal import Decimal
import openpyxl
import pytest
from rest_framework import status

from field_reports.models import (
    DailyReport,
    DailyReportActivity,
    ReportShift,
    ReportStatus,
)
from field_reports.shiraz_file1_excel import ShirazFile1ExcelBuilder


@pytest.fixture
def sample_activity_data(project, user):
    r1 = DailyReport.objects.create(
        project=project,
        created_by=user,
        report_date=date(2025, 2, 1),
        shift=ReportShift.DAY,
        work_front='بلوک ۱',
        status=ReportStatus.APPROVED,
    )
    r2 = DailyReport.objects.create(
        project=project,
        created_by=user,
        report_date=date(2025, 2, 2),
        shift=ReportShift.NIGHT,
        work_front='بلوک ۲',
        status=ReportStatus.APPROVED,
    )

    a1 = DailyReportActivity.objects.create(
        report=r1,
        source_discipline='civil',
        activity_description='قالب‌بندی ستون‌های طبقه اول',
        shift='morning',
        subcontractor_name='پیمانکاری رضایی',
        headcount=6,
        parcel='قطعه شمالی',
        block='بلوک ۱',
        floor='اول',
        location_detail='محور A-C',
        quantity=Decimal('45.50'),
        unit='مترمربع',
    )
    a2 = DailyReportActivity.objects.create(
        report=r1,
        source_discipline='civil',
        activity_description='آرماتوربندی فونداسیون',
        shift='morning',
        subcontractor_name='پیمانکاری کریمی',
        headcount=8,
        parcel='قطعه شمالی',
        block='بلوک ۱',
        floor='زیرزمین',
        location_detail='بخش غربی',
        quantity=Decimal('1200.00'),
        unit='کیلوگرم',
    )
    a3 = DailyReportActivity.objects.create(
        report=r2,
        source_discipline='electrical',
        activity_description='لوله گذاری برق سقف',
        shift='evening',
        subcontractor_name='پیمانکاری برق نوین',
        headcount=4,
        parcel='قطعه جنوبی',
        block='بلوک ۲',
        floor='دوم',
        location_detail='واحد ۳ و ۴',
        quantity=Decimal('85.00'),
        unit='مترطول',
    )
    return [a1, a2, a3]


@pytest.mark.django_db
class TestShirazFile1Excel:

    def test_user_text_is_literal_and_totals_remain_formulas(self, sample_activity_data):
        sample_activity_data[0].activity_description = '=1+1'
        sample_activity_data[0].crew_name = '=2+2'
        raw, _ = ShirazFile1ExcelBuilder(sample_activity_data).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active
        assert ws['C2'].value == '=1+1' and ws['C2'].data_type == 's'
        assert ws['M2'].value == '=2+2' and ws['M2'].data_type == 's'
        assert ws['F5'].data_type == 'f'
    def test_empty_builder(self):
        builder = ShirazFile1ExcelBuilder([])
        excel_bytes, filename = builder.build()
        assert filename == 'activity-database.xlsx'

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb.active
        assert ws.title == 'ACTIVITY'
        assert ws.sheet_view.rightToLeft is True
        assert 'هیچ فعالیتی' in str(ws['A2'].value)

    def test_populated_builder(self, sample_activity_data):
        builder = ShirazFile1ExcelBuilder(sample_activity_data)
        excel_bytes, filename = builder.build()
        assert filename == 'activity-database.xlsx'

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
        ws = wb.active
        assert ws.title == 'ACTIVITY'
        assert ws.sheet_view.rightToLeft is True

        # Check headers
        expected_headers = [
            'ردیف', 'تاریخ', 'شرح فعالیت', 'شیفت', 'پیمانکار',
            'نفرات', 'قطعه', 'بلوک', 'طبقه', 'موقعیت', 'مقدار', 'واحد', 'اکیپ', 'ناحیه'
        ]
        for col_idx, expected_h in enumerate(expected_headers, start=1):
            assert ws.cell(row=1, column=col_idx).value == expected_h

        # 3 data rows + 1 header + 1 subtotal = 5 rows
        assert ws.cell(row=2, column=3).value == 'قالب‌بندی ستون‌های طبقه اول'
        assert ws.cell(row=2, column=6).value == 6
        assert ws.cell(row=2, column=12).value == 'مترمربع'

        # Row 5 should be the subtotal row
        assert ws.cell(row=5, column=1).value == 'جمع تعداد نفرات ردیف‌ها'
        assert ws.cell(row=5, column=6).value == '=SUM(F2:F4)'

    def test_derived_rows_preserve_contractor_crew_unknowns_and_location(self, project, user):
        from projects.models import Stakeholder

        report = DailyReport.objects.create(project=project, created_by=user, report_date=date(2025, 2, 3))
        stakeholder = Stakeholder.objects.create(project=project, name='پیمانکار مرجع', role='contractor', created_by=user)
        actual = DailyReportActivity.objects.create(
            report=report, shift='shift_1', activity_description='فعالیت', contractor_ref=stakeholder,
            crew_name='اکیپ مستقل', headcount=None, parcel=None, zone='ناحیه الف',
            location_detail='محور ۱', quantity=Decimal('15'), quantity_measured=False,
        )
        missing = DailyReportActivity.objects.create(
            report=report, shift='shift_1', activity_description='فاقد پیمانکار',
            responsible_name='مسئول متفاوت', headcount=0, quantity=Decimal('0'), quantity_measured=True,
        )
        data, _ = ShirazFile1ExcelBuilder([actual, missing]).build()
        sheet = openpyxl.load_workbook(io.BytesIO(data), data_only=False).active
        assert sheet['E2'].value == 'پیمانکار مرجع'
        assert sheet['M2'].value == 'اکیپ مستقل'
        assert sheet['F2'].value is None
        assert sheet['F3'].value == 0
        assert sheet['E3'].value in (None, '—')
        assert sheet['G2'].value in (None, '—')
        assert sheet['N2'].value == 'ناحیه الف'
        assert sheet['J2'].value == 'محور ۱'
        assert sheet['K2'].value == '*'
        assert sheet['K3'].value == 0
        assert 'جمع تعداد نفرات ردیف‌ها' in sheet['A4'].value
        assert 'ناقص' in sheet['A4'].value

    def test_activity_log_excel_endpoint_authorized(self, auth_client, project, sample_activity_data):
        url = f'/api/v1/projects/{project.id}/activity-log/excel/'
        resp = auth_client.get(url)

        assert resp.status_code == status.HTTP_200_OK
        assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'Content-Disposition' in resp
        assert 'attachment; filename="activity-database.xlsx"' in resp['Content-Disposition']

        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)
        ws = wb.active
        assert ws.sheet_view.rightToLeft is True
        # 3 activities in total
        assert ws.cell(row=5, column=6).value == '=SUM(F2:F4)'

    def test_activity_log_excel_endpoint_filtered(self, auth_client, project, sample_activity_data):
        # Filter by discipline 'electrical'
        url = f'/api/v1/projects/{project.id}/activity-log/excel/?discipline=electrical'
        resp = auth_client.get(url)

        assert resp.status_code == status.HTTP_200_OK
        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=False)
        ws = wb.active

        # Only 1 electrical activity -> row 2 data, row 3 subtotal =SUM(F2:F2)
        assert ws.cell(row=2, column=3).value == 'لوله گذاری برق سقف'
        assert ws.cell(row=3, column=6).value == '=SUM(F2:F2)'

    def test_activity_log_excel_endpoint_unauthenticated(self, api_client, project):
        url = f'/api/v1/projects/{project.id}/activity-log/excel/'
        resp = api_client.get(url)
        assert resp.status_code == status.HTTP_401_UNAUTHORIZED
