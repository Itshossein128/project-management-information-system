import io
from datetime import date
from decimal import Decimal
import openpyxl
import pytest
from rest_framework import status
from django.urls import reverse

from field_reports.models import (
    DailyReport,
    DailyReportActivity,
    DailyReportConcreteLog,
    DailyReportEquipment,
    DailyReportLabor,
    DailyReportMaterial,
    DailyReportMachineryRow,
    DailyReportGeneralActivity,
    LaborCategory,
    ReportShift,
    ReportStatus,
)
from field_reports.shiraz_file4_excel import (
    SHIRAZ_DIRECT_ROLES,
    SHIRAZ_INDIRECT_ROLES,
    ShirazFile4ExcelBuilder,
)


@pytest.fixture
def sample_report_with_data(project, user):
    report = DailyReport.objects.create(
        project=project,
        created_by=user,
        report_date=date(2025, 2, 2),
        shift=ReportShift.FULL,
        work_front='بلوک ۳۲ و ۳۳',
        weather_condition='clear',
        temp_min=Decimal('5.0'),
        temp_max=Decimal('18.5'),
        site_status='active',
        prepared_by=user,
        status=ReportStatus.DRAFT,
    )
    # Indirect staff entries
    DailyReportLabor.objects.create(
        report=report,
        job_title='سرپرست کارگاه',
        custom_title='سرپرست کارگاه',
        labor_category=LaborCategory.INDIRECT,
        shift_1_count=1,
        shift_2_count=0,
        shift_3_count=0,
        absence_count=0,
        total_count=1,
    )
    DailyReportLabor.objects.create(
        report=report,
        job_title='کارشناس دفتر فنی',
        custom_title='کارشناس دفتر فنی',
        labor_category=LaborCategory.INDIRECT,
        shift_1_count=2,
        shift_2_count=1,
        shift_3_count=0,
        absence_count=1,
        total_count=3,
    )

    # Direct labor entries
    DailyReportLabor.objects.create(
        report=report,
        job_title='کارگر ساده',
        custom_title='کارگر ساده',
        labor_category=LaborCategory.DIRECT,
        shift_1_count=15,
        shift_2_count=5,
        shift_3_count=0,
        absence_count=2,
        total_count=20,
    )
    DailyReportLabor.objects.create(
        report=report,
        job_title='آرماتور بند و قالب بند',
        custom_title='آرماتور بند و قالب بند',
        labor_category=LaborCategory.DIRECT,
        shift_1_count=8,
        shift_2_count=0,
        shift_3_count=0,
        absence_count=0,
        total_count=8,
    )

    # Materials
    DailyReportMaterial.objects.create(
        report=report,
        material_description='سیمان فله تیپ ۲',
        transaction_type='receipt',
        quantity=Decimal('25.5'),
        unit='تن',
        supplier='سیمان فارس',
        consumption_location='سیلو شماره ۱',
    )

    # Concrete
    DailyReportConcreteLog.objects.create(
        report=report,
        concrete_description='بتن عیار ۳۵۰',
        volume_m3=Decimal('84.0'),
        zone='قطعه ۱۰',
        block='بلوک ۳۲',
        floor='سقف سوم',
        structural_element='دال سقف',
        notes='بخش شرقی',
    )

    # Equipment
    DailyReportEquipment.objects.create(
        report=report,
        equipment_name='پمپ بتن دکل ۳۶ متری',
        shift=ReportShift.DAY,
        status='active',
        ownership_type='rented',
        productive_hours=Decimal('7.5'),
        repair_hours=Decimal('0.5'),
    )

    # Activities
    DailyReportActivity.objects.create(
        report=report,
        activity_description='قالب‌بندی و بتن‌ریزی سقف بلوک ۳۲',
        shift='shift_1',
        subcontractor_name='پیمانکاری خنکا',
        headcount=12,
        zone='۱۰',
        block='۳۲',
        floor='۳',
        quantity=Decimal('84.0'),
        unit='متر مکعب',
    )

    return report


@pytest.mark.django_db
class TestShirazFile4Excel:

    def test_user_text_is_literal_and_totals_remain_formulas(self, sample_report_with_data):
        report = sample_report_with_data
        activity = report.activities.get()
        activity.activity_description = '=1+1'
        activity.notes = '=2+2'
        activity.save()
        report.general_notes = '=3+3'
        report.save(update_fields=['general_notes'])
        DailyReportLabor.objects.create(report=report, labor_category='direct', job_title='=4+4', shift_1_count=1)
        material = report.material_entries.get()
        material.material_description = '=5+5'
        material.save(update_fields=['material_description'])
        from field_reports.models import DailyReportVisit
        DailyReportVisit.objects.create(report=report, row_number=1, text='=6+6')
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active
        found = {cell.value: cell for row in ws for cell in row if cell.value in ('=1+1', '=2+2', '=3+3', '=4+4', '=5+5', '=6+6')}
        assert len(found) == 6
        assert all(cell.data_type == 's' for cell in found.values())
        assert any(cell.data_type == 'f' and str(cell.value).startswith('=SUM(') for row in ws for cell in row)

    def test_structured_issues_hse_and_visits_are_exported_from_unit_summary(self, project, user):
        from common.reporting_units import SpecialtyUnit
        from field_reports.services.daily_report_aggregation import rebuild_daily_report_summary
        from sub_reports.models import (
            DisciplineSubReport, DisciplineSubReportIssue, DisciplineSubReportVisit,
            DisciplineSubReportHseIncident, SubReportStatus,
        )

        day = date(2025, 2, 3)
        unit = DisciplineSubReport.objects.create(
            project=project, report_date=day, unit=SpecialtyUnit.BUILDINGS,
            discipline='civil', status=SubReportStatus.SUBMITTED, created_by=user,
        )
        DisciplineSubReportIssue.objects.create(
            sub_report=unit, row_number=1, kind='barrier', text='مانع تأمین مصالح',
            owner_user=user, due_date=date(2025, 2, 8),
        )
        DisciplineSubReportHseIncident.objects.create(
            sub_report=unit, row_number=2, text='حادثه در کارگاه',
            severity='major', immediate_action='توقف و ایمن‌سازی',
        )
        DisciplineSubReportVisit.objects.create(sub_report=unit, row_number=1, text='بازدید مشاور')
        report = rebuild_daily_report_summary(project.id, day, user=user)
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active
        cells = [[cell.value for cell in row] for row in ws.iter_rows()]
        visit = next(row for row in cells if 'بازدید مشاور' in row)
        barrier = next(row for row in cells if 'مانع تأمین مصالح' in row)
        hse = next(row for row in cells if 'حادثه در کارگاه' in row)
        assert visit[1] == 'بازدید مشاور'
        assert 'Test User' in barrier and '1403/11/20' in barrier
        assert 'توقف و ایمن‌سازی' in hse and 'Major' in hse
        assert any('مرخصی' in row for row in cells)

    def test_export_preserves_real_header_and_derived_rows(self, sample_report_with_data):
        report = sample_report_with_data
        report.project.contract_number = ''
        report.project.employer = 'کارفرمای آزمایشی'
        report.project.consultant = 'مشاور آزمایشی'
        report.project.save()
        report.weather_condition = None
        report.general_notes = 'مانع آزمایشی'
        report.save()
        DailyReportLabor.objects.create(report=report, labor_category='direct', job_title='کارگر ساده', shift_1_count=4)
        DailyReportLabor.objects.create(report=report, labor_category='direct', job_title='عنوان سفارشی', shift_1_count=7)
        DailyReportMachineryRow.objects.create(report=report, equipment_type='دستگاه تجمیعی', status_total=3, status_active=2, status_under_repair=1)
        DailyReportGeneralActivity.objects.create(report=report, row_number=1, activity_description='فعالیت عمومی خاص', location='کارگاه', quantity=2, unit='عدد', notes='شرح خاص')
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active
        text = ' '.join(str(c.value) for row in ws for c in row if c.value is not None)
        assert '1402/25319/ص' not in text
        assert '۷ الی ۱۷' not in text
        assert 'نامشخص' in text
        assert 'کارفرمای آزمایشی' in text and 'مشاور آزمایشی' in text
        assert 'دستگاه تجمیعی' in text and 'فعالیت عمومی خاص' in text and 'مانع آزمایشی' in text
        assert 'جمع نفر-شیفت' in text
        assert 'جمع کل نفر-شیفت کارگاه' in text
        rows = {row[1].value: row for row in ws.iter_rows() if len(row) > 2 and row[1].value in ('کارگر ساده', 'عنوان سفارشی')}
        assert rows['کارگر ساده'][2].value == 19
        assert rows['عنوان سفارشی'][2].value == 7
        machine = next(row for row in ws.iter_rows() if len(row) > 1 and row[1].value == 'دستگاه تجمیعی')
        assert [c.value for c in machine[2:7]] == [3, 2, 0, 0, 1]

    def test_export_keeps_unmeasured_activity_distinct_from_zero(self, sample_report_with_data):
        activity = sample_report_with_data.activities.get()
        activity.quantity = Decimal('0')
        activity.quantity_measured = False
        activity.location_detail = 'ضلع شرقی'
        activity.notes = 'توضیح ردیف'
        activity.save()
        raw, _ = ShirazFile4ExcelBuilder(sample_report_with_data).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        row = next(row for row in ws.iter_rows() if len(row) > 1 and row[1].value == activity.activity_description)
        assert row[8].value == '*'
        assert row[10].value == 'ضلع شرقی'
        assert row[11].value == 'توضیح ردیف'

    def test_activity_export_keeps_contractor_and_crew(self, sample_report_with_data):
        activity = sample_report_with_data.activities.get()
        activity.subcontractor_name = 'پیمانکار نمونه'
        activity.crew_name = 'اکیپ نمونه'
        activity.save(update_fields=['subcontractor_name', 'crew_name'])
        raw, _ = ShirazFile4ExcelBuilder(sample_report_with_data).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        row = next(row for row in ws.iter_rows() if len(row) > 3 and row[1].value == activity.activity_description)
        assert 'پیمانکار: پیمانکار نمونه' in row[3].value
        assert 'اکیپ: اکیپ نمونه' in row[3].value

    def test_header_uses_project_contract(self, sample_report_with_data):
        report = sample_report_with_data
        report.project.contract_number = 'TEST-123'
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert 'TEST-123' in ws['A1'].value

    def test_unknown_site_status_and_inactive_reason_are_not_invented(self, sample_report_with_data):
        report = sample_report_with_data
        report.site_status = ''
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert ws['E3'].value == 'وضعیت کارگاه: نامشخص'
        report.site_status = 'inactive'
        report.inactivity_reason = ''
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert ws['E3'].value == 'وضعیت کارگاه: غیرفعال (دلیل ثبت نشده)'

    def test_unrecorded_labor_is_blank_but_explicit_zero_is_preserved(self, sample_report_with_data):
        DailyReportLabor.objects.create(
            report=sample_report_with_data, labor_category='direct', job_title='کارگر ماهر',
            shift_1_count=0, shift_2_count=0, shift_3_count=0, absence_count=0,
        )
        raw, _ = ShirazFile4ExcelBuilder(sample_report_with_data).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active
        rows = {row[1].value: row for row in ws.iter_rows() if len(row) > 6 and row[1].value in ('مدیرشعبه', 'کارگر ماهر')}
        assert [c.value for c in rows['مدیرشعبه'][2:7]] == [None, None, None, None, None]
        assert [c.value for c in rows['کارگر ماهر'][2:6]] == [0, 0, 0, 0]
        assert rows['کارگر ماهر'][6].value.startswith('=SUM(')

    def test_auto_summary_warns_when_completeness_unknown(self, sample_report_with_data):
        report = sample_report_with_data
        report.is_auto_summary = True
        report.save(update_fields=['is_auto_summary'])
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert 'کامل بودن گزارش قابل احراز نیست' in ws['A5'].value
        assert 'انطباق منابع با نسخه ذخیره‌شده قابل احراز نیست' in ws['A6'].value

    def test_auto_summary_warns_when_source_changed_after_snapshot(self, project, user):
        from common.reporting_units import SpecialtyUnit
        from field_reports.services.daily_report_aggregation import rebuild_daily_report_summary
        from sub_reports.models import DisciplineSubReport, DisciplineSubReportActivity, SubReportStatus

        day = date(2025, 2, 3)
        source = DisciplineSubReport.objects.create(
            project=project, report_date=day, unit=SpecialtyUnit.BUILDINGS,
            discipline='civil', status=SubReportStatus.SUBMITTED, created_by=user, updated_by=user,
        )
        activity = DisciplineSubReportActivity.objects.create(
            sub_report=source, row_number=1, shift='shift_1', crew_name='Crew',
            activity_description='پیش از تغییر', quantity=5, unit='m2',
        )
        report = rebuild_daily_report_summary(project.id, day, user=user)
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert ws['A6'].value is None

        activity.activity_description = 'پس از تغییر'
        activity.save(update_fields=['activity_description'])
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert 'منابع پس از آخرین تجمیع تغییر کرده‌اند' in ws['A6'].value
        assert report.activities.filter(activity_description='پیش از تغییر').exists()

    def test_auto_summary_distinguishes_resolved_receipts_from_other_materials(self, sample_report_with_data):
        report = sample_report_with_data
        report.is_auto_summary = True
        report.save(update_fields=['is_auto_summary'])
        resolved = report.material_entries.get()
        resolved.receipt_resolved = True
        resolved.save(update_fields=['receipt_resolved'])
        DailyReportMaterial.objects.create(report=report, material_description='مصرف نمونه', transaction_type='issue', quantity=3, unit='تن')
        DailyReportMaterial.objects.create(report=report, material_description='رسید نامشخص', transaction_type='receipt', quantity=2, unit='تن', receipt_resolved=False)
        raw, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        rows = {cell.value: cell.row for row in ws for cell in row if cell.value in ('سیمان فله تیپ ۲', 'مصرف نمونه', 'رسید نامشخص', 'سایر گردش‌ها و رسیدهای تأییدنشده (خارج از مصالح وارده)')}
        separator = rows['سایر گردش‌ها و رسیدهای تأییدنشده (خارج از مصالح وارده)']
        assert rows['سیمان فله تیپ ۲'] < separator < rows['مصرف نمونه']
        assert separator < rows['رسید نامشخص']

    def test_empty_report_keeps_sections_and_zero_formulas(self, project, user):
        report = DailyReport.objects.create(
            project=project,
            created_by=user,
            report_date=date(2025, 2, 3),
            shift=ReportShift.FULL,
            site_status='inactive',
            inactivity_reason='holiday',
            status=ReportStatus.DRAFT,
        )

        excel_bytes, _ = ShirazFile4ExcelBuilder(report).build()
        ws = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False).active
        values = [cell.value for row in ws.iter_rows() for cell in row]
        formulas = [str(value) for value in values if isinstance(value, str) and value.startswith('=SUM')]

        assert ws.sheet_view.rightToLeft is True
        assert any('هیچ فعالیت اجرایی' in str(value) for value in values)
        assert any('موردی برای بتن‌ریزی' in str(value) for value in values)
        assert any(formula.startswith('=SUM(C') for formula in formulas)
        assert any(formula.startswith('=SUM(E') for formula in formulas)

    def test_builder_rtl_and_sections(self, sample_report_with_data):
        builder = ShirazFile4ExcelBuilder(sample_report_with_data)
        excel_bytes, filename = builder.build()

        assert filename.startswith('daily-report-')
        assert filename.endswith('.xlsx')

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
        ws = wb.active

        # RTL orientation
        assert ws.sheet_view.rightToLeft is True

        # Header contains project and title
        header_val = ws['A1'].value
        assert 'گزارش روزانه کارگاه' in header_val

        # Scan text across all cells
        cell_values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
        all_text = ' '.join(str(v) for v in cell_values)

        # 34 indirect roles checked
        for role in SHIRAZ_INDIRECT_ROLES:
            assert role in all_text, f'Role {role} missing from indirect section'

        # 32 direct roles checked
        for role in SHIRAZ_DIRECT_ROLES:
            assert role in all_text, f'Role {role} missing from direct section'

        # Native Excel formulas checked
        formula_cells = [str(v) for v in cell_values if str(v).startswith('=SUM')]
        assert len(formula_cells) >= 14, 'Expected recorded row and subtotal =SUM formulas'

        # Child tables checked
        assert 'سیمان فله تیپ ۲' in all_text
        assert 'بتن عیار ۳۵۰' in all_text
        assert 'پمپ بتن دکل ۳۶ متری' in all_text
        assert 'قالب‌بندی و بتن‌ریزی سقف بلوک ۳۲' in all_text

    def test_excel_endpoint_authorized(self, auth_client, sample_report_with_data):
        url = f'/api/v1/projects/{sample_report_with_data.project_id}/daily-reports/{sample_report_with_data.id}/excel/'
        resp = auth_client.get(url)

        assert resp.status_code == status.HTTP_200_OK
        assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'Content-Disposition' in resp
        assert 'attachment; filename="daily-report-' in resp['Content-Disposition']

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert wb.active.sheet_view.rightToLeft is True

    def test_excel_endpoint_unauthenticated(self, api_client, sample_report_with_data):
        url = f'/api/v1/projects/{sample_report_with_data.project_id}/daily-reports/{sample_report_with_data.id}/excel/'
        resp = api_client.get(url)
        assert resp.status_code == status.HTTP_401_UNAUTHORIZED
