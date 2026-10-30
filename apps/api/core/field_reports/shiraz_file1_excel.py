"""Shiraz FILE 1 Master Cumulative Activities Database Excel Builder.

Generates the flat historical activity database spreadsheet matching
Shiraz template FILE 1 (ACTIVITY sheet) with RTL layout, Tahoma font,
and columns:
ردیف, تاریخ, شرح فعالیت, شیفت, پیمانکار, نفرات, قطعه, بلوک, طبقه, موقعیت, مقدار, واحد, اکیپ, ناحیه
"""
from __future__ import annotations

import io
from typing import Iterable, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common.jalali import gregorian_to_jalali
from field_reports.models import DailyReportActivity
from field_reports.excel_safety import literalize_cell

FONT_FAMILY = 'Tahoma'

STYLE_HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
STYLE_HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color='FFFFFF')

STYLE_TOTAL_FILL = PatternFill(start_color='B4C6E7', end_color='B4C6E7', fill_type='solid')
STYLE_TOTAL_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color='000000')

STYLE_DATA_FONT = Font(name=FONT_FAMILY, size=9)

THIN_BORDER = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE'),
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)


class ShirazFile1ExcelBuilder:
    def __init__(self, queryset: Iterable[DailyReportActivity], project_code_or_name: str = ''):
        self.queryset = queryset
        self.project_code_or_name = project_code_or_name
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = 'ACTIVITY'
        self.ws.sheet_view.rightToLeft = True

    def build(self) -> Tuple[bytes, str]:
        headers = [
            'ردیف', 'تاریخ', 'شرح فعالیت', 'شیفت', 'پیمانکار',
            'نفرات', 'قطعه', 'بلوک', 'طبقه', 'موقعیت', 'مقدار', 'واحد', 'اکیپ', 'ناحیه'
        ]

        # Row 1: Headers
        for col_idx, h in enumerate(headers, start=1):
            cell = self.ws.cell(row=1, column=col_idx, value=h)
            cell.font = STYLE_HEADER_FONT
            cell.fill = STYLE_HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        row = 2
        items = list(self.queryset)
        if not items:
            self.ws.merge_cells('A2:N2')
            c = self.ws.cell(row=2, column=1, value='هیچ فعالیتی با فیلترهای مشخص‌شده یافت نشد.')
            c.font = STYLE_DATA_FONT
            c.alignment = ALIGN_CENTER
            c.border = THIN_BORDER
            row = 3
        else:
            for idx, act in enumerate(items, start=1):
                report = act.report
                r_date = gregorian_to_jalali(report.report_date) if report and report.report_date else ''
                shift_display = act.get_shift_display() if hasattr(act, 'get_shift_display') else (act.shift or '—')
                contractor = (act.contractor_ref.name if act.contractor_ref_id else act.subcontractor_name) or '—'
                headcount = act.headcount
                qty_val = float(act.quantity) if act.quantity_measured and act.quantity is not None else '*'

                self.ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=2, value=r_date).alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=3, value=act.activity_description).alignment = ALIGN_RIGHT
                self.ws.cell(row=row, column=4, value=shift_display).alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=5, value=contractor).alignment = ALIGN_RIGHT
                self.ws.cell(row=row, column=6, value=headcount).alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=7, value=act.parcel or '—').alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=8, value=act.block or '—').alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=9, value=act.floor or '—').alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=10, value=act.location_detail or '—').alignment = ALIGN_RIGHT
                self.ws.cell(row=row, column=11, value=qty_val).alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=12, value=act.unit or '—').alignment = ALIGN_CENTER
                self.ws.cell(row=row, column=13, value=act.crew_name or '—').alignment = ALIGN_RIGHT
                self.ws.cell(row=row, column=14, value=act.zone or '—').alignment = ALIGN_RIGHT

                for c_i in range(1, 15):
                    cell = literalize_cell(self.ws.cell(row=row, column=c_i))
                    cell.font = STYLE_DATA_FONT
                    cell.border = THIN_BORDER
                row += 1

            # Subtotal row for headcount
            data_end = row - 1
            self.ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            missing_count = sum(act.headcount is None for act in items)
            label = 'جمع تعداد نفرات ردیف‌ها'
            if missing_count:
                label += f' (ناقص؛ تعداد ثبت‌نشده: {missing_count})'
            lbl = self.ws.cell(row=row, column=1, value=label)
            lbl.font = STYLE_TOTAL_FONT
            lbl.fill = STYLE_TOTAL_FILL
            lbl.alignment = ALIGN_RIGHT

            tot_hc = self.ws.cell(row=row, column=6, value=f'=SUM(F2:F{data_end})')
            tot_hc.font = STYLE_TOTAL_FONT
            tot_hc.fill = STYLE_TOTAL_FILL
            tot_hc.alignment = ALIGN_CENTER

            for c_i in range(1, 15):
                cell = self.ws.cell(row=row, column=c_i)
                cell.font = STYLE_TOTAL_FONT
                cell.fill = STYLE_TOTAL_FILL
                cell.border = THIN_BORDER
            row += 1

        self._auto_fit_columns()

        buf = io.BytesIO()
        self.wb.save(buf)
        buf.seek(0)

        filename = 'activity-database.xlsx'
        return buf.getvalue(), filename

    def _auto_fit_columns(self):
        ws = self.ws
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if not val_str.startswith('='):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
