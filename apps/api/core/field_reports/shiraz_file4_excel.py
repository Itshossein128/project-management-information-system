"""Shiraz FILE 4 Master Daily Report Excel Builder.

Replicates the physical and digital layout of Shiraz National Housing
FILE 4 daily reports (RTL layout, Tahoma font, 34 staff labor roles,
32 direct labor worker roles, native Excel `=SUM(...)` subtotal and grand total
formulas, incoming materials, concrete pours, machinery, and daily activities).
"""
from __future__ import annotations

import io
from collections import defaultdict
from typing import Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common.jalali import gregorian_to_jalali, persian_day_of_week
from field_reports.models import DailyReport
from field_reports.excel_safety import literalize_cell

# Fixed 34 indirect / staff roles from Shiraz FILE 4 & migration 0010
SHIRAZ_INDIRECT_ROLES = [
    'مدیرشعبه',
    'معاونت فنی شعبه',
    'سرپرست کنترل پروژه',
    'کنترل کیفیت شعبه',
    'مدیر مالی',
    'تدارکات',
    'مدیر پروژه',
    'سرپرست کارگاه',
    'معاونت کارگاه',
    'سرپرست دفتر فنی',
    'برنامه‌ریزی و کنترل پروژه',
    'کارشناس دفتر فنی',
    'سرپرست اجرا',
    'کارشناس اجرا',
    'حسابداری',
    'واحد مکانیک',
    'واحد برق',
    'اداری و مالی',
    'کارمند اداری',
    'کنترل کیفیت',
    'سرپرست ایمنی',
    'کارشناس ایمنی',
    'انباردار',
    'نقشه‌بردار',
    'نگهبان روز و شب',
    'خدمات',
    'تأسیسات',
    'مسئول ماشین‌آلات',
    'تکنسین اجرا',
    'فناوری اطلاعات',
    'سرپرست بچینگ',
    'آزمایشگاه',
    'مهمان',
    'متفرقه ستادی',
]

# Fixed 32 direct worker roles from Shiraz FILE 4 & migration 0010
SHIRAZ_DIRECT_ROLES = [
    'کارگر ساده',
    'کارگر ماهر',
    'بنا',
    'آرماتور بند و قالب بند',
    'اکیپ بتن‌ریزی',
    'جوشکار',
    'اکیپ حفاری',
    'سرویس‌کار',
    'متفرقه (تعمیرکار-برق‌کار)',
    'پیمانکار برق و تأسیسات',
    'پلاستر کار',
    'لوله‌کش تأسیسات',
    'نصاب تاور',
    'داربست‌بند',
    'نیروی انسانی ماشین‌آلات',
    'راننده بابکت',
    'راننده کامیون',
    'راننده لودر',
    'راننده بیل مکانیکی',
    'راننده تراک میکسر',
    'راننده غلتک',
    'راننده بیل بکهو',
    'راننده جرثقیل',
    'راننده تانکر آب',
    'راننده گریدر',
    'اپراتور پمپ بتن',
    'اپراتور بچینگ',
    'اپراتور باسکول',
    'کنترل‌چی عملیات خاکی',
    'راننده سواری',
    'اپراتور دستگاه حفاری',
    'اپراتور تاور',
]

FONT_FAMILY = 'Tahoma'

STYLE_HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
STYLE_HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color='FFFFFF')

STYLE_SECTION_FILL = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
STYLE_SECTION_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color='FFFFFF')

STYLE_SUBHEADER_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
STYLE_SUBHEADER_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color='1F4E79')

STYLE_TOTAL_FILL = PatternFill(start_color='B4C6E7', end_color='B4C6E7', fill_type='solid')
STYLE_TOTAL_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color='000000')

STYLE_DATA_FONT = Font(name=FONT_FAMILY, size=9)
STYLE_BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE'),
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


class ShirazFile4ExcelBuilder:
    def __init__(self, report: DailyReport):
        self.report = report
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = 'گزارش روزانه'
        self.ws.sheet_view.rightToLeft = True

    def build(self) -> Tuple[bytes, str]:
        self._write_header()
        current_row = 10
        current_row = self._write_manpower(current_row)
        current_row += 1
        current_row = self._write_materials(current_row)
        current_row += 1
        current_row = self._write_concrete(current_row)
        current_row += 1
        current_row = self._write_machinery(current_row)
        current_row += 1
        current_row = self._write_activities(current_row)
        current_row = self._write_general_activities(current_row + 1)
        current_row = self._write_visits_and_incidents(current_row + 1)
        self._write_notes(current_row + 1)

        self._auto_fit_columns()

        buf = io.BytesIO()
        self.wb.save(buf)
        buf.seek(0)

        report_date = str(self.report.report_date)
        shift = self.report.shift or 'full'
        jalali = gregorian_to_jalali(self.report.report_date) if self.report.report_date else report_date
        formatted_date = jalali.replace('/', '-') if jalali else report_date
        filename = f'daily-report-{formatted_date}-{shift}.xlsx'

        return buf.getvalue(), filename

    def _write_header(self):
        ws = self.ws
        project_name = getattr(self.report.project, 'project_name', 'پروژه ساختمانی')
        contract_number = self.report.project.contract_number or 'ثبت نشده'
        jalali_date = gregorian_to_jalali(self.report.report_date) if self.report.report_date else ''
        dow = persian_day_of_week(self.report.report_date) if self.report.report_date else ''

        # Row 1-2: Title block
        ws.merge_cells('A1:J2')
        title_cell = ws['A1']
        title_cell.value = f'گزارش روزانه کارگاه — {project_name} (شماره قرارداد: {contract_number})'
        title_cell.font = STYLE_HEADER_FONT
        title_cell.fill = STYLE_HEADER_FILL
        title_cell.alignment = ALIGN_CENTER

        # Row 3: Date, Day, Status, Weather
        ws.merge_cells('A3:B3')
        ws['A3'] = f'تاریخ: {jalali_date}'
        ws.merge_cells('C3:D3')
        ws['C3'] = f'روز: {dow}'
        ws.merge_cells('E3:F3')
        if self.report.site_status == 'active':
            status_text = 'فعال'
        elif self.report.site_status == 'inactive':
            reason = self.report.get_inactivity_reason_display()
            status_text = f'غیرفعال ({reason})' if reason else 'غیرفعال (دلیل ثبت نشده)'
        else:
            status_text = 'نامشخص'
        ws['E3'] = f'وضعیت کارگاه: {status_text}'
        ws.merge_cells('G3:H3')
        weather_text = self.report.get_weather_condition_display() if self.report.weather_condition else 'نامشخص'
        t_min = self.report.temp_min if self.report.temp_min is not None else '—'
        t_max = self.report.temp_max if self.report.temp_max is not None else '—'
        ws['G3'] = f'وضعیت جوی: {weather_text} (دما: {t_min} تا {t_max} °C)'
        ws.merge_cells('I3:J3')
        ws['I3'] = f'شیفت کاری: {self.report.get_shift_display()}'

        # Row 4: Shift hours
        ws.merge_cells('A4:C4')
        ws['A4'] = 'کارفرما: ' + (self.report.project.employer or 'ثبت نشده')
        ws.merge_cells('D4:F4')
        ws['D4'] = 'مشاور: ' + (self.report.project.consultant or 'ثبت نشده')
        ws.merge_cells('G4:J4')
        ws['G4'] = 'ساعات کاری شیفت: ثبت نشده'

        if self.report.is_auto_summary:
            from field_reports.services.daily_report_aggregation import build_aggregation_meta

            meta = build_aggregation_meta(self.report.project_id, self.report.report_date, self.report)
            completeness = meta['completeness']
            state = completeness['state']
            if state == 'complete':
                message = 'وضعیت فعلی گزارش‌های واحدها: کامل'
            elif state == 'configuration_incomplete':
                message = 'هشدار وضعیت فعلی واحدها: پیکربندی مشخص نیست؛ کامل بودن گزارش قابل احراز نیست.'
            else:
                missing = '، '.join(completeness['missing_units']) or 'نامشخص'
                message = f'هشدار وضعیت فعلی واحدها: ناقص؛ واحدهای فاقد گزارش: {missing}'
            message += f' | آخرین تجمیع: {self.report.aggregated_at.isoformat() if self.report.aggregated_at else "ثبت نشده"}'
            ws.merge_cells('A5:L5')
            ws['A5'] = message
            ws['A5'].font = STYLE_BOLD_FONT
            ws['A5'].alignment = ALIGN_RIGHT
            ws.row_dimensions[5].height = 32
            comparison = meta['snapshot_source_comparison']
            if comparison != 'matching_fingerprint':
                ws.merge_cells('A6:L6')
                ws['A6'] = (
                    'هشدار: منابع پس از آخرین تجمیع تغییر کرده‌اند؛ این خروجی نسخه ذخیره‌شده است.'
                    if comparison == 'changed'
                    else 'هشدار: انطباق منابع با نسخه ذخیره‌شده قابل احراز نیست.'
                )
                ws['A6'].font = STYLE_BOLD_FONT
                ws['A6'].alignment = ALIGN_RIGHT
                ws.row_dimensions[6].height = 32

        # Style header rows 3-4
        for r in range(3, 5):
            for col_idx in range(1, 11):
                cell = ws.cell(row=r, column=col_idx)
                cell.font = STYLE_SUBHEADER_FONT
                cell.fill = STYLE_SUBHEADER_FILL
                cell.border = THIN_BORDER
                cell.alignment = ALIGN_CENTER

    def _write_manpower(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        # Title: نیروی انسانی
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='بخش الف و ب: نیروی انسانی ستادی (غیرمستقیم) و اجرایی (مستقیم)')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        # Section A: Staff (34 roles)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='نیروی انسانی ستادی و غیرمستقیم (۳۴ ردیف استاندارد)')
        c.font = STYLE_SUBHEADER_FONT
        c.fill = STYLE_SUBHEADER_FILL
        c.alignment = ALIGN_RIGHT
        row += 1

        absence_label = 'مرخصی' if self.report.is_auto_summary else 'مرخصی / غیبت'
        headers = ['ردیف', 'عنوان شغلی', 'شیفت ۱', 'شیفت ۲', 'شیفت ۳', absence_label, 'جمع نفر-شیفت']
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        indirect_entries = self._labor_totals('indirect')

        indirect_start = row
        for idx, role in enumerate(SHIRAZ_INDIRECT_ROLES + sorted(set(indirect_entries) - set(SHIRAZ_INDIRECT_ROLES)), start=1):
            recorded = role in indirect_entries
            s1, s2, s3, absent = indirect_entries.get(role, [None, None, None, None])

            ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
            ws.cell(row=row, column=2, value=role).alignment = ALIGN_RIGHT
            ws.cell(row=row, column=3, value=s1).alignment = ALIGN_CENTER
            ws.cell(row=row, column=4, value=s2).alignment = ALIGN_CENTER
            ws.cell(row=row, column=5, value=s3).alignment = ALIGN_CENTER
            ws.cell(row=row, column=6, value=absent).alignment = ALIGN_CENTER
            # Formula: =SUM(C{row}:E{row})
            sum_cell = ws.cell(row=row, column=7, value=f'=SUM(C{row}:E{row})' if recorded else None)
            sum_cell.alignment = ALIGN_CENTER
            sum_cell.font = STYLE_BOLD_FONT

            for c_i in range(1, 8):
                cell = ws.cell(row=row, column=c_i)
                if c_i != 7:
                    literalize_cell(cell)
                    cell.font = STYLE_DATA_FONT
                cell.border = THIN_BORDER
            row += 1

        indirect_end = row - 1

        # Subtotal row for Indirect
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sub_c = ws.cell(row=row, column=1, value='جمع نیروی انسانی غیرمستقیم (ستادی) - نفر-شیفت')
        sub_c.alignment = ALIGN_RIGHT
        sub_c.font = STYLE_TOTAL_FONT
        sub_c.fill = STYLE_TOTAL_FILL

        ws.cell(row=row, column=3, value=f'=SUM(C{indirect_start}:C{indirect_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=4, value=f'=SUM(D{indirect_start}:D{indirect_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=5, value=f'=SUM(E{indirect_start}:E{indirect_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=6, value=f'=SUM(F{indirect_start}:F{indirect_end})').alignment = ALIGN_CENTER
        indirect_subtotal_cell = f'G{row}'
        ws.cell(row=row, column=7, value=f'=SUM(G{indirect_start}:G{indirect_end})').alignment = ALIGN_CENTER

        for c_i in range(1, 8):
            cell = ws.cell(row=row, column=c_i)
            cell.font = STYLE_TOTAL_FONT
            cell.fill = STYLE_TOTAL_FILL
            cell.border = THIN_BORDER
        row += 1

        # Section B: Direct (32 roles)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='نیروی انسانی مستقیم و کارگری (۳۲ ردیف استاندارد)')
        c.font = STYLE_SUBHEADER_FONT
        c.fill = STYLE_SUBHEADER_FILL
        c.alignment = ALIGN_RIGHT
        row += 1

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        direct_entries = self._labor_totals('direct')

        direct_start = row
        for idx, role in enumerate(SHIRAZ_DIRECT_ROLES + sorted(set(direct_entries) - set(SHIRAZ_DIRECT_ROLES)), start=1):
            recorded = role in direct_entries
            s1, s2, s3, absent = direct_entries.get(role, [None, None, None, None])

            ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
            ws.cell(row=row, column=2, value=role).alignment = ALIGN_RIGHT
            ws.cell(row=row, column=3, value=s1).alignment = ALIGN_CENTER
            ws.cell(row=row, column=4, value=s2).alignment = ALIGN_CENTER
            ws.cell(row=row, column=5, value=s3).alignment = ALIGN_CENTER
            ws.cell(row=row, column=6, value=absent).alignment = ALIGN_CENTER
            sum_cell = ws.cell(row=row, column=7, value=f'=SUM(C{row}:E{row})' if recorded else None)
            sum_cell.alignment = ALIGN_CENTER
            sum_cell.font = STYLE_BOLD_FONT

            for c_i in range(1, 8):
                cell = ws.cell(row=row, column=c_i)
                if c_i != 7:
                    literalize_cell(cell)
                    cell.font = STYLE_DATA_FONT
                cell.border = THIN_BORDER
            row += 1

        direct_end = row - 1

        # Subtotal row for Direct
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sub_c2 = ws.cell(row=row, column=1, value='جمع نیروی انسانی مستقیم (کارگری) - نفر-شیفت')
        sub_c2.alignment = ALIGN_RIGHT
        sub_c2.font = STYLE_TOTAL_FONT
        sub_c2.fill = STYLE_TOTAL_FILL

        ws.cell(row=row, column=3, value=f'=SUM(C{direct_start}:C{direct_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=4, value=f'=SUM(D{direct_start}:D{direct_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=5, value=f'=SUM(E{direct_start}:E{direct_end})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=6, value=f'=SUM(F{direct_start}:F{direct_end})').alignment = ALIGN_CENTER
        direct_subtotal_cell = f'G{row}'
        ws.cell(row=row, column=7, value=f'=SUM(G{direct_start}:G{direct_end})').alignment = ALIGN_CENTER

        for c_i in range(1, 8):
            cell = ws.cell(row=row, column=c_i)
            cell.font = STYLE_TOTAL_FONT
            cell.fill = STYLE_TOTAL_FILL
            cell.border = THIN_BORDER
        row += 1

        # Grand Total Row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        gt_c = ws.cell(row=row, column=1, value='جمع کل نفر-شیفت کارگاه (ستادی + کارگری)')
        gt_c.alignment = ALIGN_RIGHT
        gt_c.font = STYLE_TOTAL_FONT
        gt_c.fill = STYLE_TOTAL_FILL

        gt_val = ws.cell(row=row, column=7, value=f'=SUM({indirect_subtotal_cell},{direct_subtotal_cell})')
        gt_val.alignment = ALIGN_CENTER
        gt_val.font = STYLE_TOTAL_FONT
        gt_val.fill = STYLE_TOTAL_FILL

        for c_i in range(1, 8):
            ws.cell(row=row, column=c_i).border = THIN_BORDER
        row += 1

        return row

    def _labor_totals(self, category):
        totals = defaultdict(lambda: [0, 0, 0, None])
        for entry in self.report.labor_entries.filter(is_deleted=False, labor_category=category):
            bucket = totals[entry.custom_title or entry.job_title]
            for index, value in enumerate((entry.shift_1_count, entry.shift_2_count, entry.shift_3_count, entry.absence_count)):
                if value is not None:
                    bucket[index] = (bucket[index] or 0) + value
        return totals

    def _write_materials(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='بخش ج: مصالح وارده به کارگاه')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        headers = ['ردیف', 'شرح مصالح', 'نوع تراکنش', 'مقدار', 'واحد', 'محل مصرف / دپو', 'تأمین‌کننده']
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        all_items = list(self.report.material_entries.filter(is_deleted=False))
        items = [item for item in all_items if item.transaction_type == 'receipt' and (not self.report.is_auto_summary or item.receipt_resolved is True)]
        if not items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = ws.cell(row=row, column=1, value='موردی برای مصالح وارده در این روز ثبت نشده است.')
            cell.font = STYLE_DATA_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            row += 1
        else:
            for idx, item in enumerate(items, start=1):
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=item.material_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=3, value=item.get_transaction_type_display()).alignment = ALIGN_CENTER
                ws.cell(row=row, column=4, value=float(item.quantity) if item.quantity is not None else 0).alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=item.unit).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=item.consumption_location or '—').alignment = ALIGN_RIGHT
                supplier_val = getattr(item, 'supplier', None) or getattr(item, 'supplier_name', None) or '—'
                ws.cell(row=row, column=7, value=supplier_val).alignment = ALIGN_RIGHT

                for c_i in range(1, 8):
                    literalize_cell(ws.cell(row=row, column=c_i)).font = STYLE_DATA_FONT
                    ws.cell(row=row, column=c_i).border = THIN_BORDER
                row += 1

        other_items = [item for item in all_items if item not in items]
        if other_items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = ws.cell(row=row, column=1, value='سایر گردش‌ها و رسیدهای تأییدنشده (خارج از مصالح وارده)')
            cell.font, cell.fill, cell.alignment = STYLE_SUBHEADER_FONT, STYLE_SUBHEADER_FILL, ALIGN_RIGHT
            row += 1
            for idx, item in enumerate(other_items, 1):
                values = [idx, item.material_description, item.get_transaction_type_display(), item.quantity, item.unit, item.consumption_location or '—', item.supplier or '—']
                for col, value in enumerate(values, 1):
                    cell = literalize_cell(ws.cell(row=row, column=col, value=value))
                    cell.font, cell.border, cell.alignment = STYLE_DATA_FONT, THIN_BORDER, ALIGN_CENTER
                row += 1
        return row

    def _write_concrete(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='بخش د: احجام بتن‌ریزی')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        headers = ['ردیف', 'شرح بتن / عیار', 'حجم (متر مکعب)', 'قطعه / زون', 'بلوک', 'طبقه', 'موقعیت']
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        items = list(self.report.concrete_logs.filter(is_deleted=False))
        if not items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = ws.cell(row=row, column=1, value='موردی برای بتن‌ریزی در این روز ثبت نشده است.')
            cell.font = STYLE_DATA_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            row += 1

            # Keep the empty section computationally valid: spreadsheet viewers
            # evaluate this as zero instead of leaving the section without a
            # numeric total/formula row.
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            total_label = ws.cell(row=row, column=1, value='جمع کل حجم بتن‌ریزی (متر مکعب)')
            total_label.font = STYLE_TOTAL_FONT
            total_label.fill = STYLE_TOTAL_FILL
            total_label.alignment = ALIGN_RIGHT
            total_value = ws.cell(row=row, column=3, value=f'=SUM(C{row - 1}:C{row - 1})')
            total_value.font = STYLE_TOTAL_FONT
            total_value.fill = STYLE_TOTAL_FILL
            total_value.alignment = ALIGN_CENTER
            for c_i in range(1, 8):
                ws.cell(row=row, column=c_i).border = THIN_BORDER
                ws.cell(row=row, column=c_i).fill = STYLE_TOTAL_FILL
            row += 1
        else:
            concrete_start = row
            for idx, item in enumerate(items, start=1):
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=item.concrete_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=3, value=float(item.volume_m3) if item.volume_m3 is not None else 0).alignment = ALIGN_CENTER
                ws.cell(row=row, column=4, value=item.zone or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=item.block or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=item.floor or '—').alignment = ALIGN_CENTER
                loc_val = getattr(item, 'structural_element', None) or getattr(item, 'notes', None) or getattr(item, 'location_notes', None) or '—'
                ws.cell(row=row, column=7, value=loc_val).alignment = ALIGN_RIGHT

                for c_i in range(1, 8):
                    literalize_cell(ws.cell(row=row, column=c_i)).font = STYLE_DATA_FONT
                    ws.cell(row=row, column=c_i).border = THIN_BORDER
                row += 1

            concrete_end = row - 1
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            c_tot = ws.cell(row=row, column=1, value='جمع کل حجم بتن‌ریزی (متر مکعب)')
            c_tot.font = STYLE_TOTAL_FONT
            c_tot.fill = STYLE_TOTAL_FILL
            c_tot.alignment = ALIGN_RIGHT

            tot_val = ws.cell(row=row, column=3, value=f'=SUM(C{concrete_start}:C{concrete_end})')
            tot_val.font = STYLE_TOTAL_FONT
            tot_val.fill = STYLE_TOTAL_FILL
            tot_val.alignment = ALIGN_CENTER

            for c_i in range(1, 8):
                ws.cell(row=row, column=c_i).border = THIN_BORDER
                ws.cell(row=row, column=c_i).fill = STYLE_TOTAL_FILL
            row += 1

        return row

    def _write_machinery(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        derived = list(self.report.machinery_rows.filter(is_deleted=False))
        if derived or self.report.is_auto_summary:
            headers = ['ردیف', 'نوع دستگاه', 'تعداد کل', 'فعال', 'آماده', 'غیرفعال', 'در تعمیر', 'ساعت کارکرد', 'ساعت توقف', 'ساعت تعمیر']
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            cell = ws.cell(row=row, column=1, value='بخش هـ: ماشین‌آلات و تجهیزات کارگاه')
            cell.font, cell.fill, cell.alignment = STYLE_SECTION_FONT, STYLE_SECTION_FILL, ALIGN_CENTER
            row += 1
            for col, title in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=title)
                cell.font, cell.fill, cell.alignment, cell.border = STYLE_SUBHEADER_FONT, STYLE_SUBHEADER_FILL, ALIGN_CENTER, THIN_BORDER
            row += 1
            for idx, item in enumerate(derived, 1):
                values = [idx, item.equipment_type, item.status_total, item.status_active, item.status_ready, item.status_inactive, item.status_under_repair, item.operating_hours, item.stoppage_hours, item.repair_hours]
                for col, value in enumerate(values, 1):
                    cell = literalize_cell(ws.cell(row=row, column=col, value=value))
                    cell.font, cell.border, cell.alignment = STYLE_DATA_FONT, THIN_BORDER, ALIGN_CENTER
                row += 1
            if not derived:
                ws.cell(row=row, column=2, value='موردی برای ماشین‌آلات در این روز ثبت نشده است.')
                row += 1
            return row

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value='بخش هـ: ماشین‌آلات و تجهیزات کارگاه')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        headers = ['ردیف', 'نام دستگاه و تجهیزات', 'شیفت', 'وضعیت', 'نوع مالکیت', 'ساعات کارکرد مفید', 'ساعات تعمیرات']
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        items = list(self.report.equipment_entries.filter(is_deleted=False))
        if not items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            cell = ws.cell(row=row, column=1, value='موردی برای ماشین‌آلات در این روز ثبت نشده است.')
            cell.font = STYLE_DATA_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            row += 1
        else:
            for idx, item in enumerate(items, start=1):
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=item.equipment_name).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=3, value=item.get_shift_display()).alignment = ALIGN_CENTER
                ws.cell(row=row, column=4, value=item.get_status_display()).alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=item.get_ownership_type_display()).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=float(item.productive_hours) if item.productive_hours is not None else 0).alignment = ALIGN_CENTER
                repair_h = float(getattr(item, 'repair_hours', 0) or getattr(item, 'idle_repair_hours', 0) or 0)
                ws.cell(row=row, column=7, value=repair_h).alignment = ALIGN_CENTER

                for c_i in range(1, 8):
                    literalize_cell(ws.cell(row=row, column=c_i)).font = STYLE_DATA_FONT
                    ws.cell(row=row, column=c_i).border = THIN_BORDER
                row += 1

        return row

    def _write_activities(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
        c = ws.cell(row=row, column=1, value='بخش و: فعالیت‌های اجرایی کارگاه (منطبق بر دیتابیس فعالیت‌ها)')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        headers = ['ردیف', 'شرح فعالیت', 'شیفت', 'پیمانکار / اکیپ', 'نفرات', 'قطعه', 'بلوک', 'طبقه', 'مقدار', 'واحد', 'موقعیت', 'توضیحات']
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        items = list(self.report.activities.filter(is_deleted=False))
        if not items:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            cell = ws.cell(row=row, column=1, value='هیچ فعالیت اجرایی در این روز ثبت نشده است.')
            cell.font = STYLE_DATA_FONT
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            row += 1

            # Preserve a zero-valued native formula for an empty activity table.
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            total_label = ws.cell(row=row, column=1, value='جمع کل نفرات فعالیت‌ها')
            total_label.font = STYLE_TOTAL_FONT
            total_label.fill = STYLE_TOTAL_FILL
            total_label.alignment = ALIGN_RIGHT
            total_value = ws.cell(row=row, column=5, value=f'=SUM(E{row - 1}:E{row - 1})')
            total_value.font = STYLE_TOTAL_FONT
            total_value.fill = STYLE_TOTAL_FILL
            total_value.alignment = ALIGN_CENTER
            for c_i in range(1, 11):
                ws.cell(row=row, column=c_i).border = THIN_BORDER
                ws.cell(row=row, column=c_i).fill = STYLE_TOTAL_FILL
            row += 1
        else:
            for idx, item in enumerate(items, start=1):
                qty_display = float(item.quantity) if item.quantity is not None and item.quantity_measured else '*'
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=item.activity_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=3, value=item.get_shift_display()).alignment = ALIGN_CENTER
                contractor = item.contractor_ref.name if item.contractor_ref_id else item.subcontractor_name
                parties = []
                if contractor:
                    parties.append(f'پیمانکار: {contractor}')
                if item.crew_name:
                    parties.append(f'اکیپ: {item.crew_name}')
                if not parties and item.responsible_name:
                    parties.append(f'مسئول: {item.responsible_name}')
                ws.cell(row=row, column=4, value=' / '.join(parties) or '—').alignment = ALIGN_RIGHT
                ws.cell(row=row, column=5, value=item.headcount).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=item.parcel or item.zone or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=7, value=item.block or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=8, value=item.floor or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=9, value=qty_display).alignment = ALIGN_CENTER
                ws.cell(row=row, column=10, value=item.unit or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=11, value=item.location_detail or '—').alignment = ALIGN_RIGHT
                ws.cell(row=row, column=12, value=item.notes or '—').alignment = ALIGN_RIGHT

                for c_i in range(1, 13):
                    literalize_cell(ws.cell(row=row, column=c_i)).font = STYLE_DATA_FONT
                    ws.cell(row=row, column=c_i).border = THIN_BORDER
                row += 1

        return row

    def _write_general_activities(self, row):
        ws = self.ws
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = ws.cell(row=row, column=1, value='فعالیت‌های عمومی کارگاه')
        cell.font, cell.fill, cell.alignment = STYLE_SECTION_FONT, STYLE_SECTION_FILL, ALIGN_CENTER
        row += 1
        for col, label in enumerate(['ردیف', 'شرح فعالیت', 'محل', 'مقدار', 'واحد', 'توضیحات'], 1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.font, cell.fill, cell.border, cell.alignment = STYLE_SUBHEADER_FONT, STYLE_SUBHEADER_FILL, THIN_BORDER, ALIGN_CENTER
        row += 1
        for item in self.report.general_activities.filter(is_deleted=False).order_by('row_number', 'id'):
            for col, value in enumerate([item.row_number, item.activity_description, item.location, item.quantity, item.unit, item.notes], 1):
                cell = literalize_cell(ws.cell(row=row, column=col, value=value))
                cell.font, cell.border, cell.alignment = STYLE_DATA_FONT, THIN_BORDER, ALIGN_RIGHT
            row += 1
        return row

    def _write_notes(self, row):
        for label, value in [('توضیحات عمومی و موانع', self.report.general_notes), ('نظر مدیریت', self.report.management_opinion), ('تأثیر شرایط جوی', self.report.weather_impact_notes)]:
            self.ws.cell(row=row, column=1, value=label).font = STYLE_BOLD_FONT
            self.ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            literalize_cell(self.ws.cell(row=row, column=2, value=value or '—')).alignment = ALIGN_RIGHT
            row += 1

    def _write_visits_and_incidents(self, row):
        """Keep structured field observations separate from free-text notes."""
        ws = self.ws

        def heading(title, headers):
            nonlocal row
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            cell = ws.cell(row=row, column=1, value=title)
            cell.font, cell.fill, cell.alignment = STYLE_SECTION_FONT, STYLE_SECTION_FILL, ALIGN_CENTER
            row += 1
            for col, label in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=label)
                cell.font, cell.fill, cell.border, cell.alignment = STYLE_SUBHEADER_FONT, STYLE_SUBHEADER_FILL, THIN_BORDER, ALIGN_CENTER
            row += 1

        def data_row(values):
            nonlocal row
            for col, value in enumerate(values, 1):
                cell = literalize_cell(ws.cell(row=row, column=col, value=value))
                cell.font, cell.border, cell.alignment = STYLE_DATA_FONT, THIN_BORDER, ALIGN_RIGHT
            row += 1

        heading('بازدیدهای کارگاه', ['ردیف', 'شرح بازدید'])
        for visit in self.report.visits.filter(is_deleted=False).order_by('row_number', 'id'):
            data_row([visit.row_number, visit.text])

        heading('مسائل، موانع و حوادث ایمنی', ['ردیف', 'نوع', 'شرح', 'مسئول پیگیری', 'موعد', 'اقدام اصلاحی / فوری', 'شدت'])
        type_labels = {'barrier': 'مانع', 'issue': 'مسئله', 'safety': 'حادثه ایمنی'}
        for incident in self.report.incidents.filter(is_deleted=False).select_related('follow_up_owner_user').order_by('row_number', 'id'):
            owner = incident.follow_up_owner_name
            if not owner and incident.follow_up_owner_user:
                owner = incident.follow_up_owner_user.get_full_name() or incident.follow_up_owner_user.username
            due = gregorian_to_jalali(incident.due_date) if incident.due_date else '—'
            data_row([
                incident.row_number, type_labels.get(incident.incident_type, incident.get_incident_type_display()),
                incident.description, owner or '—', due, incident.corrective_action or '—',
                incident.get_severity_display() if incident.severity else '—',
            ])
        return row

    def _auto_fit_columns(self):
        ws = self.ws
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if not val_str.startswith('='):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)
