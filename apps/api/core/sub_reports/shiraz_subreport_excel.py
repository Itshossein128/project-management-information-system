"""Shiraz Discipline Sub-Report Excel Builder (FILE 5 Civil / FILE 6 Electrical).

Generates trade-specific sub-report spreadsheets adhering to Shiraz templates:
- FILE 5 (Civil / Buildings):
  Includes columns: ردیف, شیفت, پیمانکار/اکیپ, استادکار, کارگر, قطعه, بلوک, طبقه, شرح فعالیت, واحد, مقدار, درصد اجرا (formatted 0.0%).
- FILE 6 (Electrical):
  Includes columns: ردیف, شیفت, نام اکیپ/پیمانکار, استادکار, کارگر, بلوک, طبقه, شرح فعالیت, واحد, مقدار (strictly omits قطعه and درصد اجرا).
"""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common.jalali import gregorian_to_jalali, persian_day_of_week
from common.reporting_units import SpecialtyUnit
from sub_reports.models import DisciplineSubReport

FONT_FAMILY = 'Tahoma'

STYLE_HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
STYLE_HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color='FFFFFF')

STYLE_SECTION_FILL = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
STYLE_SECTION_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color='FFFFFF')

STYLE_SUBHEADER_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
STYLE_SUBHEADER_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color='1F4E79')

STYLE_TOTAL_FILL = PatternFill(start_color='B4C6E7', end_color='B4C6E7', fill_type='solid')
STYLE_TOTAL_FONT = Font(name=FONT_FAMILY, size=9, bold=True, color='000000')

STYLE_DATA_FONT = Font(name=FONT_FAMILY, size=9)
STYLE_BOLD_FONT = Font(name=FONT_FAMILY, size=9, bold=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE'),
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)


class ShirazSubReportExcelBuilder:
    def __init__(self, sub_report: DisciplineSubReport):
        self.sub_report = sub_report
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.sheet_view.rightToLeft = True

        unit = getattr(sub_report, 'unit', None)
        discipline = getattr(sub_report, 'discipline', '') or ''
        self.is_civil = unit == SpecialtyUnit.BUILDINGS or discipline.lower() in ('civil', 'finishing')
        self.is_electrical = unit == SpecialtyUnit.ELECTRICAL or discipline.lower() == 'electrical'

        if self.is_civil:
            self.ws.title = 'گزارش ابنیه'
            self.discipline_name = 'اجرای ابنیه'
            self.file_prefix = 'civil-subreport'
        elif self.is_electrical:
            self.ws.title = 'گزارش تأسیسات برقی'
            self.discipline_name = 'اجرای تأسیسات برقی'
            self.file_prefix = 'electrical-subreport'
        else:
            self.ws.title = 'گزارش روزانه تخصصی'
            self.discipline_name = f'اجرای {sub_report.get_unit_display() if hasattr(sub_report, "get_unit_display") else (unit or discipline)}'
            self.file_prefix = f'{unit or discipline}-subreport'

    def build(self) -> Tuple[bytes, str]:
        self._write_header()
        current_row = 7
        current_row = self._write_activities(current_row)
        self._auto_fit_columns()

        buf = io.BytesIO()
        self.wb.save(buf)
        buf.seek(0)

        report_date = str(self.sub_report.report_date)
        jalali = gregorian_to_jalali(self.sub_report.report_date) if self.sub_report.report_date else report_date
        formatted_date = jalali.replace('/', '-') if jalali else report_date
        filename = f'{self.file_prefix}-{formatted_date}.xlsx'

        return buf.getvalue(), filename

    def _write_header(self):
        ws = self.ws
        project_name = getattr(self.sub_report.project, 'project_name', 'پروژه مسکن ملی')
        jalali_date = gregorian_to_jalali(self.sub_report.report_date) if self.sub_report.report_date else ''
        dow = persian_day_of_week(self.sub_report.report_date) if self.sub_report.report_date else ''
        form_code = getattr(self.sub_report, 'form_code', '') or 'P-REP-01'
        revision = getattr(self.sub_report, 'revision_number', '') or '00'

        col_span = 12 if self.is_civil else (10 if self.is_electrical else 11)
        last_col = get_column_letter(col_span)

        # Row 1-2: Company & Header
        ws.merge_cells(f'A1:{last_col}2')
        header_cell = ws['A1']
        header_cell.value = f'شرکت پی‌چین — گزارش روزانه واحدهای اجرا ({self.discipline_name}) — {project_name}'
        header_cell.font = STYLE_HEADER_FONT
        header_cell.fill = STYLE_HEADER_FILL
        header_cell.alignment = ALIGN_CENTER

        # Row 3: Metadata
        ws.merge_cells('A3:C3')
        ws['A3'] = f'تاریخ: {jalali_date} ({dow})'
        ws.merge_cells('D3:F3')
        weather_text = self.sub_report.get_weather_condition_display() if hasattr(self.sub_report, 'get_weather_condition_display') and self.sub_report.weather_condition else 'صاف'
        ws['D3'] = f'وضعیت هوا: {weather_text}'
        ws.merge_cells('G3:I3')
        ws['G3'] = f'کد فرم: {form_code}'
        ws.merge_cells(f'J3:{last_col}3')
        ws['J3'] = f'شماره بازنگری: {revision}'

        for c_i in range(1, col_span + 1):
            cell = ws.cell(row=3, column=c_i)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER

    def _write_activities(self, start_row: int) -> int:
        ws = self.ws
        row = start_row

        if self.is_civil:
            headers = [
                'ردیف', 'شیفت', 'نام پیمانکار / اکیپ', 'استادکار', 'کارگر',
                'قطعه', 'بلوک', 'طبقه', 'شرح فعالیت', 'واحد', 'مقدار', 'درصد اجرا'
            ]
        elif self.is_electrical:
            headers = [
                'ردیف', 'شیفت', 'نام اکیپ / پیمانکار', 'استادکار', 'کارگر',
                'بلوک', 'طبقه', 'شرح فعالیت', 'واحد', 'مقدار'
            ]
        else:
            headers = [
                'ردیف', 'شیفت', 'نام اکیپ / پیمانکار', 'استادکار', 'کارگر',
                'قطعه', 'بلوک', 'طبقه', 'شرح فعالیت', 'واحد', 'مقدار'
            ]

        col_span = len(headers)
        last_col = get_column_letter(col_span)

        # Section header
        ws.merge_cells(f'A{row}:{last_col}{row}')
        c = ws.cell(row=row, column=1, value='جدول فعالیت‌های اجرایی روزانه')
        c.font = STYLE_SECTION_FONT
        c.fill = STYLE_SECTION_FILL
        c.alignment = ALIGN_CENTER
        row += 1

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = STYLE_SUBHEADER_FONT
            cell.fill = STYLE_SUBHEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        activities = list(self.sub_report.activities.filter(is_deleted=False).order_by('row_number', 'id'))
        if not activities:
            ws.merge_cells(f'A{row}:{last_col}{row}')
            empty_cell = ws.cell(row=row, column=1, value='هیچ فعالیتی برای این گزارش ثبت نشده است.')
            empty_cell.font = STYLE_DATA_FONT
            empty_cell.alignment = ALIGN_CENTER
            empty_cell.border = THIN_BORDER
            row += 1
            return row

        data_start_row = row
        for idx, act in enumerate(activities, start=1):
            shift_display = act.shift or '۱'
            if shift_display == '1':
                shift_display = 'شیفت ۱'
            elif shift_display == '2':
                shift_display = 'شیفت ۲'
            elif shift_display == '3':
                shift_display = 'شیفت ۳'

            foreman = act.foreman_count if act.foreman_count is not None else 0
            worker = act.worker_count if act.worker_count is not None else 0
            qty = float(act.quantity) if act.quantity is not None else 0

            if self.is_civil:
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=shift_display).alignment = ALIGN_CENTER
                ws.cell(row=row, column=3, value=act.crew_name or 'نفر شرکت').alignment = ALIGN_RIGHT
                ws.cell(row=row, column=4, value=foreman).alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=worker).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=act.parcel or act.zone or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=7, value=act.block or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=8, value=act.floor or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=9, value=act.activity_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=10, value=act.unit or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=11, value=qty).alignment = ALIGN_CENTER

                # Percentage formatting 0.0%
                pct_cell = ws.cell(row=row, column=12)
                # Keep the FILE 5 column numeric even when the source is null;
                # this makes Excel's native percentage formatting stable for
                # both missing and explicit zero progress.
                pct_val = float(act.execution_percentage or 0)
                if pct_val > 1.0:
                    pct_val /= 100.0
                pct_cell.value = pct_val
                pct_cell.number_format = '0.0%'
                pct_cell.alignment = ALIGN_CENTER
            elif self.is_electrical:
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=act.shift or 1).alignment = ALIGN_CENTER
                ws.cell(row=row, column=3, value=act.crew_name or 'نفر شرکت').alignment = ALIGN_RIGHT
                ws.cell(row=row, column=4, value=foreman).alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=worker).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=act.block or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=7, value=act.floor or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=8, value=act.activity_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=9, value=act.unit or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=10, value=qty).alignment = ALIGN_CENTER
            else:
                ws.cell(row=row, column=1, value=idx).alignment = ALIGN_CENTER
                ws.cell(row=row, column=2, value=shift_display).alignment = ALIGN_CENTER
                ws.cell(row=row, column=3, value=act.crew_name or 'نفر شرکت').alignment = ALIGN_RIGHT
                ws.cell(row=row, column=4, value=foreman).alignment = ALIGN_CENTER
                ws.cell(row=row, column=5, value=worker).alignment = ALIGN_CENTER
                ws.cell(row=row, column=6, value=act.parcel or act.zone or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=7, value=act.block or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=8, value=act.floor or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=9, value=act.activity_description).alignment = ALIGN_RIGHT
                ws.cell(row=row, column=10, value=act.unit or '—').alignment = ALIGN_CENTER
                ws.cell(row=row, column=11, value=qty).alignment = ALIGN_CENTER

            for c_i in range(1, col_span + 1):
                cell = ws.cell(row=row, column=c_i)
                cell.font = STYLE_DATA_FONT
                cell.border = THIN_BORDER
            row += 1

        data_end_row = row - 1

        # Summary Subtotal Row for Foreman & Laborers
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sub_label = ws.cell(row=row, column=1, value='مجموع نفرات')
        sub_label.alignment = ALIGN_RIGHT
        sub_label.font = STYLE_TOTAL_FONT
        sub_label.fill = STYLE_TOTAL_FILL

        ws.cell(row=row, column=4, value=f'=SUM(D{data_start_row}:D{data_end_row})').alignment = ALIGN_CENTER
        ws.cell(row=row, column=5, value=f'=SUM(E{data_start_row}:E{data_end_row})').alignment = ALIGN_CENTER

        for c_i in range(1, col_span + 1):
            cell = ws.cell(row=row, column=c_i)
            cell.font = STYLE_TOTAL_FONT
            cell.fill = STYLE_TOTAL_FILL
            cell.border = THIN_BORDER
        row += 1

        return row

    def _auto_fit_columns(self):
        ws = self.ws
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = str(val)
                    if not val_str.startswith('='):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)
