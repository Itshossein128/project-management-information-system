"""
Comprehensive automated tests for Feature 042: Department Activity Reports & Excel Workbooks.
Covers:
- executed_quantity decimal support for non-warehouse departments.
- civil <-> buildings slug alias interchangeability.
- Single-day report (.xlsx) generation with project header, RTL, and dynamic =SUM() formula.
- Weekly report (8 sheets) with 7 daily tabs + 1 Summary tab and native =SUM() formulas.
- Filtered RTL export with Persian headers and executed_quantity.
- Atomic batch import rollback on invalid rows (0 rows committed) with structured Persian error messages.
- Permission elevation enforcing edit_reports for batch import (HTTP 403 when lacking).
"""
import io
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from business_meta.models import Project
from inventory.models import Department, DepartmentActivityRecord
from inventory.serializers import DepartmentActivityRecordSerializer
from master_data.models import ProjectMember, ProjectMemberRole, Role, RolePermission
from projects.models import Stakeholder
from projects.tests.contractor_select_helpers import make_active_contractor

User = get_user_model()


class DepartmentActivityReportsTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='report_lead', password='password', mobile='09121112233'
        )
        self.user.is_superuser = True
        self.user.save()

        self.project = Project.objects.create(
            project_code='PRJ-042',
            project_name='برج سامان البرز',
            employer='شرکت عمران کاران',
            start_date=date(2026, 1, 1),
        )

        self.contractor_a = make_active_contractor(self.project, self.user, name='پارس سازه')
        self.contractor_b = make_active_contractor(self.project, self.user, name='نوین بتن')

        self.client.force_authenticate(user=self.user)

    def test_executed_quantity_model_and_serializer(self):
        """FR-001: Model & serializer persist positive decimal executed_quantity with 3 decimals."""
        record = DepartmentActivityRecord.objects.create(
            project=self.project,
            department='buildings',
            date=date(2026, 10, 17),
            location='بلوک شرقی',
            activity_description='بتن‌ریزی فونداسیون',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('45.500'),
            unit='مترمکعب',
        )
        self.assertEqual(record.executed_quantity, Decimal('45.500'))

        # Serializer validation
        serializer = DepartmentActivityRecordSerializer(record)
        self.assertEqual(Decimal(str(serializer.data['executed_quantity'])), Decimal('45.500'))

        # Rejection of negative executed_quantity
        invalid_data = {
            'department': 'buildings',
            'date': '2026-10-17',
            'location': 'بلوک شرقی',
            'activity_description': 'بتن‌ریزی فونداسیون',
            'contractor': self.contractor_a.name,
            'contractor_ref': str(self.contractor_a.id),
            'executed_quantity': '-12.000',
            'unit': 'مترمکعب',
        }
        val_serializer = DepartmentActivityRecordSerializer(
            data=invalid_data, context={'project_id': self.project.id}
        )
        self.assertFalse(val_serializer.is_valid())
        self.assertIn('executed_quantity', val_serializer.errors)

    def test_civil_and_buildings_slug_alias_interchangeability(self):
        """FR-002: civil and buildings slugs are accepted interchangeably with 200 OK."""
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='buildings',
            date=date(2026, 10, 17),
            location='سازه اصلی',
            activity_description='قالب‌بندی',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('100.000'),
            unit='مترمربع',
        )

        daily_url = reverse(
            'project-department-activity-record-daily-report',
            kwargs={'project_pk': self.project.id},
        )

        # Query using department=civil
        resp_civil = self.client.get(daily_url, {'department': 'civil', 'date': '2026-10-17'})
        self.assertEqual(resp_civil.status_code, status.HTTP_200_OK)
        self.assertEqual(
            resp_civil['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        # Query using department=buildings
        resp_buildings = self.client.get(
            daily_url, {'department': 'buildings', 'date': '2026-10-17'}
        )
        self.assertEqual(resp_buildings.status_code, status.HTTP_200_OK)

        # Export using civil
        export_url = reverse(
            'project-department-activity-record-export',
            kwargs={'project_pk': self.project.id},
        )
        resp_export = self.client.get(export_url, {'department': 'civil'})
        self.assertEqual(resp_export.status_code, status.HTTP_200_OK)

    def test_daily_excel_report_rtl_and_dynamic_sum_formula(self):
        """FR-003, FR-004, FR-005, FR-010: Daily report .xlsx has project header, RTL, and =SUM()."""
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='electrical',
            date=date(2026, 10, 17),
            location='طبقه ۳',
            activity_description='سیم‌کشی روشنایی',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('250.000'),
            unit='متر',
        )
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='electrical',
            date=date(2026, 10, 17),
            location='طبقه ۴',
            activity_description='لوله‌گذاری برق',
            contractor=self.contractor_b.name,
            contractor_ref=self.contractor_b,
            executed_quantity=Decimal('150.000'),
            unit='متر',
        )

        url = reverse(
            'project-department-activity-record-daily-report',
            kwargs={'project_pk': self.project.id},
        )
        response = self.client.get(url, {'department': 'electrical', 'date': '2026-10-17'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)

        # Title block in row 1
        self.assertIn('گزارش روزانه فعالیت دپارتمان', str(ws.cell(row=1, column=1).value))
        self.assertIn(self.project.project_name, str(ws.cell(row=1, column=1).value))

        # Check total row with native =SUM formula
        # Row 4 is header, row 5 is item 1, row 6 is item 2, row 7 is total
        total_cell = ws.cell(row=7, column=6)
        self.assertEqual(total_cell.value, '=SUM(F5:F6)')

    def test_weekly_excel_report_8_sheets_and_sum_formulas(self):
        """FR-006, FR-007, FR-008: Weekly report has exactly 8 sheets (7 daily + 1 Summary) with =SUM()."""
        week_start = date(2026, 10, 17)  # Saturday
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='mechanical',
            date=week_start,  # شنبه
            location='موتورخانه',
            activity_description='نصب پمپ سیرکولاتور',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('2.000'),
            unit='دستگاه',
        )
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='mechanical',
            date=week_start + timedelta(days=2),  # دوشنبه
            location='رایزر شرقی',
            activity_description='لوله‌کشی آبگرم',
            contractor=self.contractor_b.name,
            contractor_ref=self.contractor_b,
            executed_quantity=Decimal('40.000'),
            unit='متر',
        )

        url = reverse(
            'project-department-activity-record-weekly-report',
            kwargs={'project_pk': self.project.id},
        )
        response = self.client.get(url, {'department': 'mechanical', 'week_start': '2026-10-17'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        wb = load_workbook(io.BytesIO(response.content))
        expected_sheets = [
            'شنبه',
            'یکشنبه',
            'دوشنبه',
            'سه‌شنبه',
            'چهارشنبه',
            'پنجشنبه',
            'جمعه',
            'خلاصه هفتگی',
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        for sheet_name in expected_sheets:
            sheet = wb[sheet_name]
            self.assertTrue(sheet.sheet_view.rightToLeft)

        # Check summary sheet formulas
        summary_ws = wb['خلاصه هفتگی']
        self.assertIn('گزارش جمع‌بندی هفتگی', str(summary_ws.cell(row=1, column=1).value))
        # Total row in summary sheet has =SUM()
        last_row = summary_ws.max_row
        weekly_total_cell = summary_ws.cell(row=last_row, column=12)
        self.assertTrue(str(weekly_total_cell.value).startswith('=SUM('))

    def test_empty_week_spanning_month_boundary_keeps_formula_totals(self):
        """US2/AC4: empty days and an empty month-crossing week remain valid workbooks."""
        url = reverse(
            'project-department-activity-record-weekly-report',
            kwargs={'project_pk': self.project.id},
        )
        response = self.client.get(
            url, {'department': 'civil', 'week_start': '2026-10-31'}
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        wb = load_workbook(io.BytesIO(response.content), data_only=False)
        self.assertEqual(len(wb.sheetnames), 8)
        for sheet_name in wb.sheetnames[:7]:
            ws = wb[sheet_name]
            self.assertIn('رکوردی در این روز ثبت نشده است', str(ws['A5'].value))
            self.assertTrue(str(ws['F6'].value).startswith('=SUM('))
        summary = wb['خلاصه هفتگی']
        self.assertIn('رکوردی در این هفته ثبت نشده است', str(summary['A5'].value))
        self.assertTrue(str(summary['E6'].value).startswith('=SUM('))
        self.assertTrue(str(summary['L6'].value).startswith('=SUM('))

    def test_civil_alias_is_consistent_across_list_reports_export_and_import(self):
        """FR-002: civil/buildings alias is accepted by every activity surface."""
        record = DepartmentActivityRecord.objects.create(
            project=self.project,
            department='buildings',
            date=date(2026, 10, 31),
            location='بلوک شرقی',
            activity_description='بتن‌ریزی',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('5.500'),
            unit='مترمکعب',
        )
        list_url = reverse(
            'project-department-activity-record-list',
            kwargs={'project_pk': self.project.id},
        )
        list_response = self.client.get(list_url, {'department': 'civil'})
        self.assertEqual(list_response.status_code, status.HTTP_200_OK)
        self.assertEqual(list_response.data['count'], 1)

        daily_url = reverse(
            'project-department-activity-record-daily-report',
            kwargs={'project_pk': self.project.id},
        )
        weekly_url = reverse(
            'project-department-activity-record-weekly-report',
            kwargs={'project_pk': self.project.id},
        )
        export_url = reverse(
            'project-department-activity-record-export',
            kwargs={'project_pk': self.project.id},
        )
        for endpoint, params in (
            (daily_url, {'department': 'civil', 'date': '2026-10-31'}),
            (weekly_url, {'department': 'civil', 'week_start': '2026-10-31'}),
            (export_url, {'department': 'civil'}),
        ):
            self.assertEqual(self.client.get(endpoint, params).status_code, status.HTTP_200_OK)

        wb = Workbook()
        ws = wb.active
        ws.append(('تاریخ', 'موقعیت', 'شرح فعالیت', 'پیمانکار', 'مقدار اجرا', 'واحد', 'توضیحات'))
        ws.append(('2026-11-01', 'بلوک غربی', 'آرماتوربندی', self.contractor_a.name, '۱۲.۵۰۰', 'کیلوگرم', ''))
        # Arabic-Indic numerals are a separate input alphabet from Persian numerals.
        ws.append(('2026-11-02', 'بلوک غربی', 'قالب‌بندی', self.contractor_a.name, '٢٠.٢٥٠', 'مترمربع', ''))
        buf = io.BytesIO()
        wb.save(buf)
        import_url = reverse(
            'project-department-activity-record-import',
            kwargs={'project_pk': self.project.id},
        )
        upload = SimpleUploadedFile('alias_digits.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        import_response = self.client.post(
            import_url + '?department=civil', {'file': upload}, format='multipart'
        )
        self.assertEqual(import_response.status_code, status.HTTP_200_OK)
        imported = DepartmentActivityRecord.objects.filter(
            project=self.project, department='buildings'
        ).order_by('date')
        self.assertEqual(imported.count(), 3)
        self.assertEqual(imported.get(date=date(2026, 11, 1)).executed_quantity, Decimal('12.500'))
        self.assertEqual(imported.get(date=date(2026, 11, 2)).executed_quantity, Decimal('20.250'))

    def test_filtered_export_rtl_and_persian_headers(self):
        """FR-009, FR-010: Filtered export produces RTL .xlsx with Persian column headers."""
        DepartmentActivityRecord.objects.create(
            project=self.project,
            department='machinery',
            date=date(2026, 10, 15),
            location='محوطه',
            activity_description='خاکبرداری گود',
            contractor=self.contractor_a.name,
            contractor_ref=self.contractor_a,
            executed_quantity=Decimal('350.000'),
            unit='مترمکعب',
        )

        url = reverse(
            'project-department-activity-record-export',
            kwargs={'project_pk': self.project.id},
        )
        response = self.client.get(
            url, {'department': 'machinery', 'contractor': 'پارس سازه'}
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            (
                'تاریخ',
                'موقعیت / جبهه کاری',
                'شرح فعالیت',
                'پیمانکار',
                'مقدار اجرا',
                'واحد',
                'توضیحات',
            ),
        )
        self.assertEqual(float(rows[1][4]), 350.0)

    def test_atomic_batch_import_rollback_on_invalid_row(self):
        """FR-012, FR-013, FR-014: Non-numeric quantity causes full rollback (0 rows) and Persian error."""
        pre_count = DepartmentActivityRecord.objects.filter(project=self.project).count()

        wb = Workbook()
        ws = wb.active
        ws.append(('تاریخ', 'موقعیت', 'شرح فعالیت', 'پیمانکار', 'مقدار اجرا', 'واحد', 'توضیحات'))
        ws.append(('2026-10-10', 'موقعیت ۱', 'آیتم ۱', 'پارس سازه', '10.000', 'متر', 'توضیح'))
        ws.append(('2026-10-11', 'موقعیت ۲', 'آیتم ۲', 'پارس سازه', '20.000', 'متر', 'توضیح'))
        # Row 4 has intentional invalid string in executed_quantity:
        ws.append(('2026-10-12', 'موقعیت ۳', 'آیتم ۳', 'پارس سازه', 'ده متر', 'متر', 'توضیح'))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        url = reverse(
            'project-department-activity-record-import',
            kwargs={'project_pk': self.project.id},
        )
        file_obj = SimpleUploadedFile(
            'import_activities.xlsx',
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        response = self.client.post(
            url + '?department=civil', {'file': file_obj}, format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('detail', response.data)
        self.assertIn('سطر ۴', response.data['detail'])
        self.assertIn('مقدار اجرا باید عددی باشد', response.data['detail'])

        # Strict All-or-Nothing check: 0 rows saved
        post_count = DepartmentActivityRecord.objects.filter(project=self.project).count()
        self.assertEqual(post_count, pre_count)

    def test_atomic_batch_import_invalid_contractor_rollback(self):
        """FR-014: Row with nonexistent contractor triggers atomic rollback and Persian error."""
        wb = Workbook()
        ws = wb.active
        ws.append(('تاریخ', 'موقعیت', 'شرح فعالیت', 'پیمانکار', 'مقدار اجرا', 'واحد', 'توضیحات'))
        ws.append(('2026-10-10', 'موقعیت ۱', 'آیتم ۱', 'پیمانکار ناموجود در این پروژه', '10.000', 'متر', ''))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        url = reverse(
            'project-department-activity-record-import',
            kwargs={'project_pk': self.project.id},
        )
        file_obj = SimpleUploadedFile('import_invalid_con.xlsx', buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        response = self.client.post(url + '?department=civil', {'file': file_obj}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('پیمانکار مشخص شده در این پروژه معتبر نیست', response.data['detail'])
        self.assertEqual(DepartmentActivityRecord.objects.filter(project=self.project).count(), 0)

    def test_atomic_batch_import_success(self):
        """FR-011, FR-013: Valid Excel sheet persists all rows atomically."""
        wb = Workbook()
        ws = wb.active
        ws.append(('تاریخ', 'موقعیت', 'شرح فعالیت', 'پیمانکار', 'مقدار اجرا', 'واحد', 'توضیحات'))
        ws.append(('2026-10-10', 'بلوک A', 'آرماتوربندی', 'پارس سازه', '12.500', 'کیلوگرم', 'تایید شد'))
        ws.append(('2026-10-11', 'بلوک B', 'قالب‌بندی', 'پارس سازه', '45.000', 'مترمربع', 'تایید شد'))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        url = reverse(
            'project-department-activity-record-import',
            kwargs={'project_pk': self.project.id},
        )
        file_obj = SimpleUploadedFile('import_success.xlsx', buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        response = self.client.post(url + '?department=civil', {'file': file_obj}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['created'], 2)
        self.assertEqual(DepartmentActivityRecord.objects.filter(project=self.project).count(), 2)

    def test_batch_import_permission_elevation(self):
        """FR-015: User with only view_reports is denied with HTTP 403 Forbidden."""
        # Create user with only view_reports role
        viewer_user = User.objects.create_user(
            username='viewer_user', password='password', mobile='09129998877'
        )
        role_viewer, _ = Role.objects.get_or_create(role_name='Viewer Role')
        RolePermission.objects.get_or_create(role=role_viewer, permission_codename='view_reports')

        member = ProjectMember.objects.create(project=self.project, user=viewer_user, status='active')
        ProjectMemberRole.objects.create(member=member, role=role_viewer)

        client_viewer = APIClient()
        client_viewer.force_authenticate(user=viewer_user)

        wb = Workbook()
        ws = wb.active
        ws.append(('تاریخ', 'موقعیت', 'شرح فعالیت', 'پیمانکار', 'مقدار اجرا', 'واحد', 'توضیحات'))
        ws.append(('2026-10-10', 'موقعیت', 'شرح فعالیت', 'پارس سازه', '10.000', 'متر', ''))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        url = reverse(
            'project-department-activity-record-import',
            kwargs={'project_pk': self.project.id},
        )
        file_obj = SimpleUploadedFile('test.xlsx', buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # Attempt import with viewer account
        response = client_viewer.post(url + '?department=civil', {'file': file_obj}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
